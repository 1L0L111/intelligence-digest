from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.models import DigestEvent, UserProfile

def export_docx(events:list[DigestEvent],profile:UserProfile,path:str|Path)->Path:
    doc=Document(); sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.75)
    styles=doc.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10)
    h=doc.add_heading('INTELLIGENCE DIGEST',0); h.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f'Роль: {profile.role}  |  Интересы: {", ".join(profile.interests)}')
    for i,e in enumerate(sorted(events,key=lambda x:x.score,reverse=True),1):
        doc.add_heading(f'{i}. {e.title}',level=1)
        meta=doc.add_paragraph(); meta.add_run(f'{e.category} · relevance {e.score:.0%}').bold=True
        doc.add_paragraph(e.summary)
        if e.why_it_matters: doc.add_paragraph('Почему важно: '+e.why_it_matters)
        if e.details:
            details=', '.join(f'{k}: {v}' for k,v in e.details.items() if v not in (None,[],{},False,''))
            if details: doc.add_paragraph('Детали: '+details)
        doc.add_paragraph('Причины отбора: '+(', '.join(e.reasons) or 'semantic relevance'))
        doc.add_paragraph('Источники:',style=None)
        for s in e.sources: doc.add_paragraph(s,style='List Bullet')
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True); doc.save(path); return path
